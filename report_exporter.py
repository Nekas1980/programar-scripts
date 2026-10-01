from pathlib import Path
from datetime import datetime
from html import escape



def prepare_report_folder(base_dir):

    report_dir = Path(base_dir) / "reports"
    report_dir.mkdir(exist_ok=True)

    return report_dir


def report_filename(language, extension):

    now = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

    return f"security_report_{language}_{now}.{extension}"


def translate_result(result, msg):

    risk = msg.get(
        result["risk"],
        result["risk"].upper()
    )

    event = msg.get(
        result["event"],
        result["event"]
    )

    if result["night"]:
        night = msg.get("yes", "YES")
    else:
        night = msg.get("no", "NO")

    return risk, event, night


# ============================================================
# TXT
# ============================================================

def export_txt(results, msg, report_dir, language):

    filename = report_filename(
        language,
        "txt"
    )

    report_path = report_dir / filename

    with open(
        report_path,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(
            msg.get(
                "title",
                "SECURITY LOG ANALYSIS"
            )
        )

        file.write("\n")
        file.write("=" * 90)
        file.write("\n\n")

        file.write(
            f"{msg.get('generated_at', 'Generated at')}: "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        )

        file.write(
            f"{msg.get('total_logs', 'Total logs')}: "
            f"{len(results)}\n\n"
        )


        for number, result in enumerate(
            results,
            start=1
        ):

            risk, event, night = (
                translate_result(
                    result,
                    msg
                )
            )

            file.write(
                f"--- {msg.get('event', 'Event')} "
                f"{number} ---\n"
            )

            file.write(
                f"{msg.get('timestamp', 'Timestamp')}: "
                f"{result['timestamp']}\n"
            )

            file.write(
                f"{msg.get('severity', 'Severity')}: "
                f"{result['severity']}\n"
            )

            file.write(
                f"{msg.get('risk', 'Risk')}: "
                f"{risk}\n"
            )

            file.write(
                f"{msg.get('event', 'Event')}: "
                f"{event}\n"
            )

            file.write(
                f"{msg.get('ip', 'IP Address')}: "
                f"{result['ip']}\n"
            )

            file.write(
                f"{msg.get('email', 'Email')}: "
                f"{result['email']}\n"
            )

            file.write(
                f"{msg.get('night', 'Night')}: "
                f"{night}\n"
            )

            file.write(
                f"HTTP: "
                f"{result['http_code']}\n"
            )

            file.write(
                f"{msg.get('full_line', 'Full Log Line')}:\n"
                f"{result['full_line']}\n"
            )

            file.write("\n")


    return report_path


# ============================================================
# HTML
# ============================================================

def export_html(results, msg, report_dir, language):

    filename = report_filename(
        language,
        "html"
    )

    report_path = report_dir / filename


    rows = []


    for result in results:

        risk, event, night = (
            translate_result(
                result,
                msg
            )
        )

        rows.append(
            f"""
            <tr>
                <td>{escape(str(result['timestamp']))}</td>
                <td>{escape(str(result['severity']))}</td>
                <td>{escape(str(risk))}</td>
                <td>{escape(str(event))}</td>
                <td>{escape(str(result['ip']))}</td>
                <td>{escape(str(result['email']))}</td>
                <td>{escape(str(night))}</td>
                <td>{escape(str(result['http_code']))}</td>
                <td>{escape(str(result['full_line']))}</td>
            </tr>
            """
        )


    html = f"""
<!DOCTYPE html>

<html lang="{escape(language)}">

<head>

<meta charset="UTF-8">

<title>
{escape(msg.get("title", "Security Log Analysis"))}
</title>

<style>

body {{
    font-family: Arial, sans-serif;
    margin: 30px;
}}

table {{
    border-collapse: collapse;
    width: 100%;
}}

th, td {{
    border: 1px solid #999;
    padding: 8px;
    text-align: left;
    vertical-align: top;
}}

th {{
    background: #eeeeee;
}}

</style>

</head>

<body>

<h1>
{escape(msg.get("title", "Security Log Analysis"))}
</h1>

<p>
{escape(msg.get("generated_at", "Generated at"))}:
{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
</p>

<p>
{escape(msg.get("total_logs", "Total logs"))}:
{len(results)}
</p>

<table>

<thead>

<tr>
    <th>{escape(msg.get("timestamp", "Timestamp"))}</th>
    <th>{escape(msg.get("severity", "Severity"))}</th>
    <th>{escape(msg.get("risk", "Risk"))}</th>
    <th>{escape(msg.get("event", "Event"))}</th>
    <th>{escape(msg.get("ip", "IP Address"))}</th>
    <th>{escape(msg.get("email", "Email"))}</th>
    <th>{escape(msg.get("night", "Night"))}</th>
    <th>HTTP</th>
    <th>{escape(msg.get("full_line", "Full Log Line"))}</th>
</tr>

</thead>

<tbody>

{''.join(rows)}

</tbody>

</table>

</body>

</html>
"""


    with open(
        report_path,
        "w",
        encoding="utf-8"
    ) as file:

        file.write(html)


    return report_path


# ============================================================
# DOCX - MICROSOFT WORD
# ============================================================

def export_docx(results, msg, report_dir, language):

    try:

        from docx import Document

    except ImportError:

        return None


    filename = report_filename(
        language,
        "docx"
    )

    report_path = report_dir / filename


    document = Document()


    document.add_heading(
        msg.get(
            "title",
            "Security Log Analysis"
        ),
        level=1
    )


    document.add_paragraph(
        f"{msg.get('generated_at', 'Generated at')}: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    )


    document.add_paragraph(
        f"{msg.get('total_logs', 'Total logs')}: "
        f"{len(results)}"
    )


    table = document.add_table(
        rows=1,
        cols=9
    )


    headers = [
        msg.get("timestamp", "Timestamp"),
        msg.get("severity", "Severity"),
        msg.get("risk", "Risk"),
        msg.get("event", "Event"),
        msg.get("ip", "IP Address"),
        msg.get("email", "Email"),
        msg.get("night", "Night"),
        "HTTP",
        msg.get(
            "full_line",
            "Full Log Line"
        )
    ]


    for index, header in enumerate(headers):

        table.rows[0].cells[index].text = (
            str(header)
        )


    for result in results:

        risk, event, night = (
            translate_result(
                result,
                msg
            )
        )


        row = table.add_row().cells


        values = [
            result["timestamp"],
            result["severity"],
            risk,
            event,
            result["ip"],
            result["email"],
            night,
            result["http_code"],
            result["full_line"]
        ]


        for index, value in enumerate(values):

            row[index].text = str(value)


    document.save(report_path)

    return report_path


# ============================================================
# ODT - LIBREOFFICE
# ============================================================

def export_odt(results, msg, report_dir, language):

    try:

        from odf.opendocument import OpenDocumentText
        from odf.text import P
        from odf.table import (
            Table,
            TableRow,
            TableCell
        )

    except ImportError:

        return None


    filename = report_filename(
        language,
        "odt"
    )

    report_path = report_dir / filename


    document = OpenDocumentText()


    document.text.addElement(
        P(
            text=msg.get(
                "title",
                "Security Log Analysis"
            )
        )
    )


    document.text.addElement(
        P(
            text=(
                f"{msg.get('generated_at', 'Generated at')}: "
                f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            )
        )
    )


    table = Table()


    headers = [
        msg.get("timestamp", "Timestamp"),
        msg.get("severity", "Severity"),
        msg.get("risk", "Risk"),
        msg.get("event", "Event"),
        msg.get("ip", "IP Address"),
        msg.get("email", "Email"),
        msg.get("night", "Night"),
        "HTTP",
        msg.get(
            "full_line",
            "Full Log Line"
        )
    ]


    header_row = TableRow()


    for header in headers:

        cell = TableCell()

        cell.addElement(
            P(
                text=str(header)
            )
        )

        header_row.addElement(cell)


    table.addElement(header_row)


    for result in results:

        risk, event, night = (
            translate_result(
                result,
                msg
            )
        )


        values = [
            result["timestamp"],
            result["severity"],
            risk,
            event,
            result["ip"],
            result["email"],
            night,
            result["http_code"],
            result["full_line"]
        ]


        table_row = TableRow()


        for value in values:

            cell = TableCell()

            cell.addElement(
                P(
                    text=str(value)
                )
            )

            table_row.addElement(cell)


        table.addElement(table_row)


    document.text.addElement(table)

    document.save(str(report_path))

    return report_path


        # ========================================================
    # EXPORT REPORTS
    # ========================================================

    report_dir = prepare_report_folder(BASE_DIR)

    print()
    print("=== EXPORTAR RELATÓRIO ===")
    print("1 - TXT")
    print("2 - HTML")
    print("3 - DOCX / Word")
    print("4 - ODT / LibreOffice")
    print("5 - TODOS")

    report_option = input("Escolha uma opção: ").strip()

    generated_reports = []

    if report_option == "1" or report_option == "5":
        path = export_txt(results, msg, report_dir, language)
        generated_reports.append(path)

    if report_option == "2" or report_option == "5":
        path = export_html(results, msg, report_dir, language)
        generated_reports.append(path)

    if report_option == "3" or report_option == "5":
        path = export_docx(results, msg, report_dir, language)

        if path:
            generated_reports.append(path)
        else:
            print("DOCX indisponível. Instala: python-docx")

    if report_option == "4" or report_option == "5":
        path = export_odt(results, msg, report_dir, language)

        if path:
            generated_reports.append(path)
        else:
            print("ODT indisponível. Instala: odfpy")

    print()

    if generated_reports:
        print("=== RELATÓRIOS CRIADOS ===")

        for report in generated_reports:
            print(report)
    else:
        print("Nenhum relatório foi criado.")